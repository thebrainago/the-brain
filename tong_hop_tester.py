# -*- coding: utf-8 -*-
r"""tong_hop_tester.py - gom reports/drive_bot_tester.jsonl thanh MOT workbook reports/drive_bot_tester.xlsx.
Mot dong / lan chay (lan chay moi nhat thang lan cu cung `nhan`). Cot 'danh gia' chi la nhan so bo:
CO_LAI_DD<80 = lai > 0 va maxDD < 80% (cong chan chu du an, tai_lieu/LAN_EA_THO.md); khong phai bang chung (Model 1 la khao sat)."""
import json, pathlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

LAB = pathlib.Path(__file__).parent
rows = {}
for l in (LAB / "reports" / "drive_bot_tester.jsonl").read_text(encoding="utf-8").splitlines():
    try:
        j = json.loads(l)
    except ValueError:
        continue
    rows[j["nhan"]] = j
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "tester"
cot = ["nhan", "ea", "set", "khung", "tu", "den", "model", "von", "slot", "giay", "xong", "so_lenh", "lai_rong", "lai_pct_von",
       "dd_pct", "pf", "thang_pct", "chat_luong_pct", "danh_gia", "ghi_chu"]
ws.append(cot)
for k in sorted(rows):
    j = rows[k]
    s = j.get("so") or {}
    lenh = s.get("so_lenh") or 0
    lai = s.get("lai_rong")
    dd = s.get("dd_pct")
    if not j.get("xong"):
        dg, gc = "KHONG_CHAY", j.get("loi", "")
    elif not lenh:
        dg, gc = "0_LENH", (j.get("log_cuoi") or "")[:200]
    elif lai is not None and lai > 0 and dd is not None and dd < 80:
        dg, gc = "CO_LAI_DD<80", ""
    else:
        dg, gc = "LO_HOAC_DD>=80", ""
    ws.append([j["nhan"], j["ea"], j["set"], j["khung"], j["tu"], j["den"], j["model"], j["von"], j.get("slot", 0), j["giay"], j["xong"],
               lenh, lai, round(lai / j["von"] * 100, 1) if lai is not None else None, dd, s.get("pf"),
               (s.get("thang") or {}).get("pct"), s.get("chat_luong_pct"), dg, gc])
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDDDDD")
for r in range(2, ws.max_row + 1):
    v = ws.cell(r, cot.index("danh_gia") + 1).value
    ws.cell(r, cot.index("danh_gia") + 1).fill = PatternFill("solid", fgColor="C6EFCE" if v == "CO_LAI_DD<80" else "FFC7CE")
for i, w in enumerate([34, 34, 34, 7, 11, 11, 6, 8, 5, 7, 6, 8, 11, 10, 8, 6, 9, 9, 16, 60]):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
ws.freeze_panes = "B2"
wb.save(LAB / "reports" / "drive_bot_tester.xlsx")
print(len(rows), "dong ->", LAB / "reports" / "drive_bot_tester.xlsx")
